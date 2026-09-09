from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

path=Path(r'C:\Users\Legion 5 pro\Desktop\cyber sec\Assignment_5_Lung_Cancer_Detection_ML_Style.docx')
d=Document(path)
def add(title, body, items):
 d.add_page_break();h=d.add_paragraph(style='Heading 1');h.paragraph_format.space_after=Pt(8);r=h.add_run(title);r.font.name='Courier New';r._element.rPr.rFonts.set(qn('w:ascii'),'Courier New');r.font.size=Pt(11);r.bold=True
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(8);r=p.add_run(body);r.font.name='Courier New';r._element.rPr.rFonts.set(qn('w:ascii'),'Courier New');r.font.size=Pt(10)
 for item in items:
  p=d.add_paragraph();p.paragraph_format.space_after=Pt(6);r=p.add_run(item);r.font.name='Courier New';r._element.rPr.rFonts.set(qn('w:ascii'),'Courier New');r.font.size=Pt(10)
add('13. APPENDIX A - CNN ARCHITECTURE','This appendix summarizes the academic CNN workflow used in the demonstration.', ['Input: 128 x 128 CT image after resizing and normalization.','Feature layers: convolution and pooling identify useful image patterns.','Output: a probability score for the illustrative class label.'])
add('14. APPENDIX B - EVALUATION LIMITATIONS','This classroom model is not a clinical diagnostic tool and requires appropriate validation before any real-world use.', ['Use separate training, validation and test records.','Measure sensitivity, specificity and false-negative behavior.','Require clinician review and external validation before deployment.'])
d.save(path)
